import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .db import account_balances


def period_bounds(d_from: str, d_to: str):
    return d_from + "T00:00:00", d_to + "T23:59:59"


def presets(today: date):
    first = today.replace(day=1)
    prev_last = first - timedelta(days=1)
    return {
        "today": (today, today), "yesterday": (today - timedelta(days=1),) * 2,
        "week": (today - timedelta(days=today.weekday()), today),
        "month": (first, today), "prev_month": (prev_last.replace(day=1), prev_last),
    }


def build(con, d_from: str, d_to: str, account_id: int | None = None) -> dict:
    lo, hi = period_bounds(d_from, d_to)
    q = ("SELECT t.*, c.name cat, COALESCE(c.exclude,0) excl, a.name acc, a2.name acc2 FROM transactions t "
         "JOIN accounts a ON a.id=t.account_id LEFT JOIN accounts a2 ON a2.id=t.to_account_id "
         "LEFT JOIN categories c ON c.id=t.category_id "
         "WHERE t.deleted_at IS NULL AND t.occurred_at BETWEEN ? AND ?")
    args = [lo, hi]
    if account_id:
        q += " AND (t.account_id=? OR t.to_account_id=?)"
        args += [account_id, account_id]
    rows = con.execute(q + " ORDER BY t.occurred_at", args).fetchall()
    inc = exp = 0
    cats, accs, days, transfers, excl = {}, {}, {}, [], {}
    for t in rows:
        a = t["amount"]
        if t["kind"] == "transfer":
            transfers.append(t)
            if not account_id or t["account_id"] == account_id:
                accs.setdefault(t["acc"], [0, 0, 0])[2] -= a
            if t["acc2"] and (not account_id or t["to_account_id"] == account_id):
                accs.setdefault(t["acc2"], [0, 0, 0])[2] += a
            continue
        if t["excl"]:
            e = excl.setdefault((t["kind"], t["cat"]), [0, 0])
            e[0] += a
            e[1] += 1
            continue
        if t["kind"] == "income":
            inc += a
        else:
            exp += a
        key = (t["kind"], t["cat"] or "Без категорії")
        c = cats.setdefault(key, [0, 0])
        c[0] += a
        c[1] += 1
        ac = accs.setdefault(t["acc"], [0, 0, 0])
        ac[0 if t["kind"] == "income" else 1] += a
        d = days.setdefault(t["occurred_at"][:10], [0, 0])
        d[0 if t["kind"] == "income" else 1] += a
    by_cat = sorted(({"kind": k[0], "category": k[1], "sum": v[0], "count": v[1]} for k, v in cats.items()),
                    key=lambda x: (x["kind"] != "income", -x["sum"]))
    return {
        "from": d_from, "to": d_to, "income": inc, "expense": exp, "net": inc - exp,
        "transfers_total": sum(t["amount"] for t in transfers),
        "by_category": by_cat,
        "excluded": [{"kind": k[0], "category": k[1], "sum": v[0], "count": v[1]} for k, v in sorted(excl.items())],
        "by_account": [{"account": k, "income": v[0], "expense": v[1], "transfers_net": v[2]} for k, v in accs.items()],
        "by_day": [{"date": k, "income": v[0], "expense": v[1]} for k, v in sorted(days.items())],
        "transactions": rows,
    }


def balances(con) -> dict:
    bal = account_balances(con)
    out = []
    for a in con.execute("SELECT * FROM accounts WHERE active=1 AND kind!='credit' ORDER BY sort,id"):
        out.append({"id": a["id"], "name": a["name"], "kind": a["kind"], "source": a["source"],
                    "balance": bal[a["id"]], "bank_balance": a["bank_balance"],
                    "synced_at": a["last_sync_at"], "status": a["last_sync_status"]})
    return {"accounts": out, "total": sum(x["balance"] for x in out)}


def uah(k: int) -> float:
    return round(k / 100, 2)


KIND_UA = {"income": "Надходження", "expense": "Витрата", "transfer": "Переказ"}


def to_xlsx(rep: dict, bal: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Підсумок"
    ws.append([f"Звіт {rep['from']} — {rep['to']}"])
    ws["A1"].font = Font(bold=True, size=13)
    for lbl, v in (("Надходження", rep["income"]), ("Витрати", rep["expense"]), ("Результат", rep["net"]),
                   ("Перекази між касами (не враховані)", rep["transfers_total"])):
        ws.append([lbl, uah(v)])
    ws.append([])
    ws.append(["Рахунок", "Надходження", "Витрати", "Перекази (нетто)"])
    for a in rep["by_account"]:
        ws.append([a["account"], uah(a["income"]), uah(a["expense"]), uah(a["transfers_net"])])
    ws.append([])
    ws.append(["Залишки на зараз"])
    for a in bal["accounts"]:
        ws.append([a["name"], uah(a["balance"])])
    ws.append(["Разом", uah(bal["total"])])
    ws2 = wb.create_sheet("По категоріях")
    ws2.append(["Тип", "Категорія", "Сума", "Кількість"])
    for c in rep["by_category"]:
        ws2.append([KIND_UA[c["kind"]], c["category"], uah(c["sum"]), c["count"]])
    ws3 = wb.create_sheet("Операції")
    ws3.append(["Дата", "Рахунок", "Тип", "Категорія", "Сума", "Контрагент", "Опис", "Джерело"])
    for t in rep["transactions"]:
        acc = t["acc"] + (f" → {t['acc2']}" if t["acc2"] else "")
        ws3.append([t["occurred_at"].replace("T", " "), acc, KIND_UA[t["kind"]], t["cat"] or "",
                    uah(t["amount"]) * (1 if t["kind"] != "expense" else -1), t["counterparty"],
                    t["description"], t["source"]])
    for w in (ws, ws2, ws3):
        for i in range(1, w.max_column + 1):
            w.column_dimensions[get_column_letter(i)].width = 22
    ws3.column_dimensions["G"].width = 60
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
