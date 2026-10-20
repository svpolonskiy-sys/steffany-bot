"""Payroll from the salon's monthly Google Sheet (actuals) + a rate-based estimate for months not filled yet.

Sheet rules (as in the sheet's own formulas): hairdressers, nail masters, podologist 40% of the month's till,
brow master 60%, cosmetologist 20% laser / 35% cosmetology / 50% injections (net of drug cost) , +10% of personal
retail sales; administrators = shift rate + share of sales; "other staff" fixed.
Cost of a person = Зарплата + Відпускні + Лікарняні - штрафи/списання.
"""
import io
import re
from datetime import date, timedelta

import requests
from openpyxl import load_workbook

from . import config
from .categorize import norm
from .db import get_setting, now_local, set_setting

SECTIONS = [("адміністратор", "Адміністратори"), ("перукар", "Перукарі"), ("майстри манікюр", "Майстри манікюру"),
            ("подолог", "Подолог"), ("бровіст", "Бровісти"), ("косметолог", "Косметолог"),
            ("інші працівник", "Інші працівники")]
SECTION_ORDER = [lbl for _, lbl in SECTIONS]
MONTHS = {"січ": 1, "лют": 2, "берез": 3, "квіт": 4, "трав": 5, "черв": 6, "лип": 7, "серп": 8, "верес": 9,
          "жовт": 10, "листоп": 11, "груд": 12}


class PayrollError(Exception):
    pass


def fetch(sheet_id: str, session=requests) -> bytes:
    r = session.get(f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx", timeout=90)
    if r.status_code != 200 or r.content[:2] != b"PK":
        raise PayrollError("Не вдалось завантажити таблицю зарплати (перевірте доступ за посиланням)")
    return r.content


def _month(title: str, default_year: int):
    t = norm(title)
    for stem, m in MONTHS.items():
        if stem in t:
            y = re.search(r"(20\d\d)", t)
            return f"{int(y.group(1)) if y else default_year}-{m:02d}"
    return None


def _num(v) -> int:
    try:
        return int(round(float(v) * 100))
    except (TypeError, ValueError):
        return 0


def parse(content: bytes, default_year: int | None = None) -> dict[str, list[dict]]:
    default_year = default_year or date.today().year
    wb = load_workbook(io.BytesIO(content), data_only=True)
    out: dict[str, list[dict]] = {}
    for ws in wb.worksheets:
        ym = _month(ws.title, default_year)
        if not ym:
            continue
        section, hdr, rows = None, None, []
        for row in ws.iter_rows(values_only=True):
            a = row[0] if row else None
            if isinstance(a, str) and a.strip():
                na = norm(a)
                if na.startswith("№"):
                    hdr = {norm(str(h).replace("\n", " ")): i for i, h in enumerate(row) if h}
                    continue
                if na.startswith("зведен"):
                    section = None
                    continue
                for key, label in SECTIONS:
                    if na.startswith(key):
                        section, hdr = label, None
                        break
                continue
            if not (section and hdr and isinstance(a, (int, float)) and len(row) > 1 and isinstance(row[1], str)):
                continue
            name = row[1].strip()
            if not name or norm(name) == "разом":
                continue
            get = lambda *keys: next((_num(row[hdr[k]]) for k in keys if k in hdr and hdr[k] < len(row)), 0)
            fines = next((_num(row[i]) for k, i in hdr.items() if k.startswith("штраф") and i < len(row)), 0)
            gross = get("зарплата", "зп")
            vac, sick = get("відпускні"), get("лікарняні")
            if gross or vac or sick:
                num = lambda k: float(row[hdr[k]] or 0) if k in hdr and hdr[k] < len(row) and isinstance(row[hdr[k]], (int, float)) else 0.0
                rows.append(dict(section=section, name=name, gross=gross, vacation=vac, sick=sick, fines=abs(fines),
                                 taxes=get("податки"), sales=get("каса за місяць"), hours=num("відпрац. годин"),
                                 shifts=num("відпрац. змін"), rate=get("ставка за годину"), retail=get("сума продажів"),
                                 adv_card=get("аванс ф-1"), sal_card=get("зп 1-форма"),
                                 adv_cash=get("аванс готівка") or (get("аванс") if section == "Інші працівники" else 0),
                                 to_pay=get("до виплати")))
        if rows:
            out[ym] = rows
    return out


def sync(con, session=requests) -> list[str]:
    """Reload every month that has numbers; empty months never wipe stored data."""
    if not config.PAYROLL_SHEET_ID:
        raise PayrollError("Не задано PAYROLL_SHEET_ID")
    data = parse(fetch(config.PAYROLL_SHEET_ID, session))
    for ym, rows in data.items():
        con.execute("DELETE FROM payroll WHERE month=?", (ym,))
        for r in rows:
            con.execute("INSERT INTO payroll(month,section,name,gross,vacation,sick,fines,taxes,sales,hours,shifts,rate,retail,"
                        "adv_card,sal_card,adv_cash,to_pay) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        (ym, r["section"], r["name"], r["gross"], r["vacation"], r["sick"], r["fines"], r["taxes"], r["sales"],
                         r["hours"], r["shifts"], r["rate"], r["retail"], r["adv_card"], r["sal_card"], r["adv_cash"], r["to_pay"]))
    set_setting(con, "payroll_synced_at", now_local())
    set_setting(con, "payroll_error", "")
    con.commit()
    return sorted(data)


def maybe_sync(con, every_hours: float = 3):
    if not config.PAYROLL_SHEET_ID:
        return
    last = get_setting(con, "payroll_synced_at")
    if last:
        from datetime import datetime
        if (datetime.fromisoformat(now_local()) - datetime.fromisoformat(last)).total_seconds() < every_hours * 3600:
            return
    try:
        sync(con)
    except Exception as e:
        set_setting(con, "payroll_error", f"{type(e).__name__}: {e}"[:250])


MASTER_SECTIONS = ("Перукарі", "Майстри манікюру")
# ledger categories that are settlements of the sheet payroll (cleaners etc. under "Витрати на персонал" are not in the sheet)
SHEET_CATEGORIES = ("Зарплата", "Премія", "Відпускні", "ЛІКАРНЯНІ")


def month_costs(con, ym: str) -> dict[str, int] | None:
    """Actual payroll of a month. A month counts as filled only when masters are in it (a half-filled current month
    with just a few advances must not be mistaken for the whole payroll)."""
    rows = con.execute("SELECT section, SUM(gross+vacation+sick-fines) c FROM payroll WHERE month=? GROUP BY section", (ym,)).fetchall()
    d = {r["section"]: r["c"] for r in rows}
    if sum(d.get(s, 0) for s in MASTER_SECTIONS) <= 0:
        return None
    return d


def covered_months(con) -> list[str]:
    return [r["month"] for r in con.execute("SELECT DISTINCT month FROM payroll ORDER BY month") if month_costs(con, r["month"])]


def _section_of_role(role: str) -> str:
    n = norm(role)
    if "brow" in role.lower() or "бров" in n:
        return "Бровісти"
    if "косметолог" in n:
        return "Косметолог"
    if "нігт" in n:
        return "Майстри манікюру"
    if "подолог" in n:
        return "Подолог"
    if "перукар" in n:
        return "Перукарі"
    return "Інші працівники"


# cosmetologist effective shares of Alteg price (calibrated on the sheet: cosmetology 35% of net ~75%, injections
# 50% of net ~72% after drug cost, laser 20%)
COSM_RATE = {"Лазерна епіляція": 0.20, "Косметологія": 0.35 * 0.75, "Ін'єкційна косметологія": 0.50 * 0.72}


def estimate(con, lo: date, hi: date, ref_month: str | None) -> dict[str, int]:
    """Accrued payroll for a period not yet in the sheet, using the sheet's own rules:
      masters: 40% of their till (brow 60%, cosmetologist 20/35/50%), +10% of personal retail sales;
      administrators: hourly rate x hours + 4% of the salon's services pool (split by shifts) + 10% of personal retail;
      Мирослава (manager): 1.5% of the salon's services; other staff: fixed amounts of the latest filled month."""
    a, b = lo.isoformat() + "T00:00:00", hi.isoformat() + "T23:59:59"
    days = (hi - lo).days + 1
    dim = ((lo.replace(day=28) + timedelta(days=4)).replace(day=1) - lo.replace(day=1)).days
    res: dict[str, float] = {}
    services = 0
    for r in con.execute("SELECT staff_role role, category_title cat, SUM(amount) s FROM alteg_pay "
                         "WHERE occurred_at BETWEEN ? AND ? AND staff_id IS NOT NULL AND category_title!='Сертифікати' "
                         "GROUP BY 1,2", (a, b)):
        sec = _section_of_role(r["role"])
        n = norm(r["role"])
        if n.startswith("корпоративна") or "адміністратор" in n:
            continue
        if r["cat"] == "Товари":
            rate = 0.10
        else:
            services += r["s"]
            rate = COSM_RATE.get(r["cat"], 0.30) if sec == "Косметолог" else 0.60 if sec == "Бровісти" else 0.40
        res[sec] = res.get(sec, 0) + r["s"] * rate
    if ref_month:
        ref = con.execute("SELECT * FROM payroll WHERE month=?", (ref_month,)).fetchall()
        y, m = map(int, ref_month.split("-"))
        ref_dim = (date(y + (m == 12), m % 12 + 1, 1) - date(y, m, 1)).days
        adm = [r for r in ref if r["section"] == "Адміністратори"]
        if adm:
            rate_h = max((r["rate"] for r in adm), default=7000)
            hours_per_day = sum(r["hours"] for r in adm) / ref_dim
            ref_services = sum(r["sales"] for r in ref if r["section"] in ("Перукарі", "Майстри манікюру", "Подолог", "Бровісти")) \
                + sum(r["sales"] for r in ref if r["section"] == "Косметолог")
            retail_ratio = (sum(r["retail"] for r in adm) * 0.10 / ref_services) if ref_services else 0
            res["Адміністратори"] = rate_h * hours_per_day * days + 0.04 * services + retail_ratio * services
        for r in ref:
            if r["section"] != "Інші працівники":
                continue
            amt = (0.015 * services) if "банцер" in norm(r["name"]) else r["gross"] * days / dim
            res["Інші працівники"] = res.get("Інші працівники", 0) + amt
    return {k: int(round(v)) for k, v in res.items() if v}


def for_period(con, lo: date, hi: date) -> tuple[dict[str, int], list[str], dict[str, dict[str, int]]]:
    """Payroll accrued in [lo, hi] by section and the months that are estimates."""
    hist = date.fromisoformat(config.HISTORY_START)
    total: dict[str, int] = {}
    est: list[str] = []
    per_month: dict[str, dict[str, int]] = {}
    covered = covered_months(con)
    m = lo.replace(day=1)
    while m <= hi:
        nxt = (m.replace(day=28) + timedelta(days=4)).replace(day=1)
        a_, b_ = max(m, lo), min(nxt - timedelta(days=1), hi)
        ym = m.strftime("%Y-%m")
        if b_ >= hist:
            a_ = max(a_, hist)
            dim = (nxt - m).days
            costs = month_costs(con, ym)
            if costs:
                f = ((b_ - a_).days + 1) / dim
                for sec, v in costs.items():
                    total[sec] = total.get(sec, 0) + int(round(v * f))
                    per_month.setdefault(ym, {})[sec] = int(round(v * f))
            else:
                ref = max((c for c in covered if c < ym), default=None)
                e = estimate(con, a_, b_, ref)
                for sec, v in e.items():
                    total[sec] = total.get(sec, 0) + v
                    per_month.setdefault(ym, {})[sec] = v
                if any(e.values()):
                    est.append(ym)
        m = nxt
    return total, est, per_month


def compensation(con, lo: date, hi: date) -> int:
    """Taxes (ПДФО + ВЗ) that masters compensate to the company: the 'Податки' column of the sheet, pro-rated."""
    total = 0.0
    for ym in covered_months(con):
        y, m = map(int, ym.split("-"))
        ms = date(y, m, 1)
        nxt = (ms.replace(day=28) + timedelta(days=4)).replace(day=1)
        a_, b_ = max(ms, lo), min(nxt - timedelta(days=1), hi)
        if b_ < a_:
            continue
        t = con.execute("SELECT COALESCE(SUM(taxes),0) s FROM payroll WHERE month=?", (ym,)).fetchone()["s"]
        total += t * ((b_ - a_).days + 1) / (nxt - ms).days
    return int(round(total))


def payment_plan(con, ym: str) -> dict:
    """What the sheet says is paid in each wave for the month worked: card advance (~15-19th of the month),
    card salary 'Ф-1' (~1-3rd of next month), cash part from safe/cash desk (7th of next month)."""
    r = con.execute("SELECT COALESCE(SUM(adv_card),0) a, COALESCE(SUM(sal_card),0) f, COALESCE(SUM(adv_cash),0) c, "
                    "COALESCE(SUM(to_pay),0) p, COALESCE(SUM(gross+vacation+sick-fines),0) cost, COALESCE(SUM(taxes),0) tx "
                    "FROM payroll WHERE month=?", (ym,)).fetchone()
    return dict(adv_card=r["a"], sal_card=r["f"], cash=r["c"] + r["p"], cost=r["cost"], compensation=r["tx"])


def payments_recorded(con, ym: str) -> dict:
    """Ledger payments tied to the month they pay for (not the month they were made in)."""
    from .categorize import norm as _n
    from .pl import accrual_date, group_of
    y, m = map(int, ym.split("-"))
    ms = date(y, m, 1)
    nxt = (ms.replace(day=28) + timedelta(days=4)).replace(day=1)
    out = {"adv_card": 0, "sal_card": 0, "cash": 0, "items": []}
    q = ("SELECT t.id, t.amount, t.occurred_at, t.description, a.kind akind, a.name aname, COALESCE(c.name,'') cname, "
         "COALESCE(c.exclude,0) ex, COALESCE(c.pl_group,'') pg FROM transactions t JOIN accounts a ON a.id=t.account_id "
         "LEFT JOIN categories c ON c.id=t.category_id WHERE t.deleted_at IS NULL AND t.kind='expense' AND t.occurred_at BETWEEN ? AND ?")
    for t in con.execute(q, (ms.isoformat() + "T00:00:00", (nxt + timedelta(days=12)).isoformat() + "T23:59:59")):
        if group_of(t["cname"] or "Інші витрати", "expense", t["ex"], t["pg"]) != "labour" or t["cname"] not in SHEET_CATEGORIES:
            continue
        paid = date.fromisoformat(t["occurred_at"][:10])
        txt = _n(t["description"])
        if t["akind"] == "fop":
            if accrual_date(t["description"], paid).strftime("%Y-%m") != ym:
                continue
            kind = "adv_card" if "аванс" in txt else "sal_card"
        elif t["akind"] == "cash":
            if not (ms <= paid <= nxt + timedelta(days=12)):
                continue
            kind = "cash"
        else:
            continue
        out[kind] += t["amount"]
        out["items"].append(dict(kind=kind, at=t["occurred_at"][:10], amount=t["amount"], account=t["aname"], text=t["description"][:60]))
    return out


HISTORY_START = date(2026, 9, 1)   # the ledger starts here: waves paid earlier were never booked


def control(con, today: date) -> list[dict]:
    """Standing check: for every month with a sheet, each payment wave that is already due must match what was
    actually paid in the ledger. Returns the waves that do not (plan - recorded, in kopecks)."""
    out = []
    for ym in covered_months(con):
        y, m = map(int, ym.split("-"))
        nxt = (date(y, m, 28) + timedelta(days=4)).replace(day=1)
        plan, rec = payment_plan(con, ym), payments_recorded(con, ym)
        for key, label, due in (("adv_card", "Аванс (Ф-1)", date(y, m, 20)),
                                ("sal_card", "ЗП безготівка (Ф-1)", nxt.replace(day=4)),
                                ("cash", "ЗП готівкою", nxt.replace(day=8))):
            if today >= due and due >= HISTORY_START and abs(plan[key] - rec[key]) >= 100:
                out.append(dict(month=ym, wave=label, plan=plan[key], rec=rec[key], diff=plan[key] - rec[key]))
    return out
