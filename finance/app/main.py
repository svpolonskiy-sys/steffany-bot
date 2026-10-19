import asyncio
import csv
import io
import logging
import re
import time
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta
from urllib.parse import urlparse

from fastapi import Depends, FastAPI, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware

from . import categorize as cat, config, payroll, pl, privat, reports, sync
from .db import (account_balances, add_audit, connect, get_setting, insert_tx, lock, now_local, set_setting,
                 update_tx)
from .security import encrypt, hash_password, verify_password

log = logging.getLogger("app")
BASE_DIR = __file__.rsplit("/", 1)[0]
ROLES = {"owner": "Власник", "accountant": "Бухгалтер", "admin": "Адміністратор"}
KINDS = {"income": "Надходження", "expense": "Витрата", "transfer": "Переказ"}
ACC_KINDS = {"fop": "Рахунок ФОП", "cash": "Готівкова каса", "card": "Приватна картка", "credit": "Кредитна лінія"}
SRC = {"manual": "вручну", "privat": "Приватбанк", "alteg": "Alteg"}

con = None  # set in create_app / lifespan


def get_con():
    global con
    if con is None:
        con = connect()
    return con


def bootstrap(c):
    if not c.execute("SELECT 1 FROM users").fetchone() and config.OWNER_EMAIL and config.OWNER_PASSWORD:
        c.execute("INSERT INTO users(email,name,role,pw_hash) VALUES(?,?,?,?)",
                  (config.OWNER_EMAIL.lower(), "Власник", "owner", hash_password(config.OWNER_PASSWORD)))
    if not c.execute("SELECT 1 FROM accounts").fetchone():
        rows = [("ФОП 1", "fop", "privat"), ("ФОП 2", "fop", "privat"), ("Каса салон", "cash", "alteg"),
                ("Сейф", "cash", "manual"), ("Робочі гроші", "card", "manual")]
        for i, (n, k, s) in enumerate(rows):
            c.execute("INSERT INTO accounts(name,kind,source,sort) VALUES(?,?,?,?)", (n, k, s, i))
        sid = c.execute("SELECT id FROM accounts WHERE name='Сейф'").fetchone()["id"]
        set_setting(c, "safe_account_id", str(sid))
    cat.seed(c)
    cat.migrate(c)
    c.commit()


async def scheduler():
    while True:
        await asyncio.sleep(60)
        try:
            c = get_con()
            last = get_setting(c, "last_sync_all")
            due = not last or (datetime.now(config.KYIV).replace(tzinfo=None)
                               - datetime.fromisoformat(last)).total_seconds() >= config.SYNC_INTERVAL_SEC
            if due and sync.in_work_window():
                await asyncio.to_thread(sync.sync_all, c)
        except Exception:
            log.exception("scheduler error")


@asynccontextmanager
async def lifespan(app):
    bootstrap(get_con())
    task = asyncio.create_task(scheduler()) if config.SCHEDULER_ENABLED else None
    yield
    if task:
        task.cancel()


app = FastAPI(lifespan=lifespan, docs_url=None, redoc_url=None, openapi_url=None)
app.add_middleware(SessionMiddleware, secret_key=config.SECRET_KEY, max_age=60 * 60 * 24 * 30,
                   same_site="lax", https_only=config.COOKIE_SECURE, session_cookie="sf_session")
app.mount("/static", StaticFiles(directory=f"{BASE_DIR}/static"), name="static")
tpl = Jinja2Templates(directory=f"{BASE_DIR}/templates")


def money(k, sign=False):
    if k is None:
        return "—"
    s = f"{abs(k) / 100:,.2f}".replace(",", " ").replace(".", ",")
    pre = "-" if k < 0 else ("+" if sign and k > 0 else "")
    return f"{pre}{s} ₴"


tpl.env.filters["money"] = money
tpl.env.filters["dt"] = lambda s: (s or "").replace("T", " ")[:16]
tpl.env.globals.update(suggest_pattern=cat.suggest_pattern, pl_group_of=lambda c: pl.group_of(c["name"], c["kind"], c["exclude"], c["pl_group"]), PL_EXP=pl.EXPENSE_GROUPS, PL_INC=pl.INCOME_GROUPS, ROLES=ROLES, KINDS=KINDS, ACC_KINDS=ACC_KINDS, SRC=SRC)


class Redirect(Exception):
    def __init__(self, url):
        self.url = url


@app.exception_handler(Redirect)
async def _redir(request, exc):
    return RedirectResponse(exc.url, status_code=303)


@app.middleware("http")
async def csrf_origin(request: Request, call_next):
    if request.method == "POST":
        o = request.headers.get("origin") or request.headers.get("referer")
        if o and urlparse(o).netloc != request.headers.get("host"):
            return Response("Forbidden", status_code=403)
    return await call_next(request)


def current_user(request: Request):
    uid = request.session.get("uid")
    u = get_con().execute("SELECT * FROM users WHERE id=? AND active=1", (uid,)).fetchone() if uid else None
    if not u:
        raise Redirect("/login")
    return u


def owner_only(u=Depends(current_user)):
    if u["role"] != "owner":
        raise HTTPException(403, "Лише для власника")
    return u


def staff_cat(u=Depends(current_user)):
    if u["role"] not in ("owner", "accountant"):
        raise HTTPException(403, "Лише для власника та бухгалтера")
    return u


def api_user(request: Request):
    tok = request.headers.get("x-api-token", "")
    if config.API_TOKEN and tok and tok == config.API_TOKEN:
        return True
    return current_user(request)


def render(request, name, user=None, **ctx):
    flash = request.session.pop("flash", None)
    return tpl.TemplateResponse(request, name, {"user": user, "flash": flash, "now": now_local(), **ctx})


def flash(request, msg):
    request.session["flash"] = msg


def parse_amount(s: str) -> int:
    s = re.sub(r"[\s ₴грнUAH]", "", s or "", flags=re.I).replace(",", ".")
    try:
        v = round(float(s) * 100)
    except ValueError:
        raise HTTPException(400, "Некоректна сума")
    if v <= 0:
        raise HTTPException(400, "Сума має бути більшою за нуль")
    return v


def parse_dt(s: str) -> str:
    try:
        return datetime.fromisoformat(s).strftime("%Y-%m-%dT%H:%M:%S")
    except ValueError:
        raise HTTPException(400, "Некоректна дата")


# ---------- auth ----------
_fail: dict[str, list[float]] = {}


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return render(request, "login.html")


@app.post("/login")
def login(request: Request, email: str = Form(...), password: str = Form(...)):
    ip = request.client.host if request.client else "?"
    recent = [t for t in _fail.get(ip, []) if time.time() - t < 600]
    if len(recent) >= 8:
        flash(request, "Забагато спроб. Спробуйте за кілька хвилин.")
        return RedirectResponse("/login", status_code=303)
    u = get_con().execute("SELECT * FROM users WHERE email=? AND active=1", (email.strip().lower(),)).fetchone()
    if not u or not verify_password(password, u["pw_hash"]):
        _fail[ip] = recent + [time.time()]
        flash(request, "Невірний email або пароль")
        return RedirectResponse("/login", status_code=303)
    request.session.clear()
    request.session["uid"] = u["id"]
    return RedirectResponse("/", status_code=303)


@app.post("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


@app.get("/health")
def health():
    return {"ok": True}


# ---------- dashboard ----------
@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, u=Depends(current_user)):
    c = get_con()
    b = reports.balances(c)
    today = datetime.now(config.KYIV).date()
    day = reports.build(c, today.isoformat(), today.isoformat())
    takings = reports.daily_cash(c, today.isoformat(), today.isoformat())["days"][0]
    month = reports.build(c, today.replace(day=1).isoformat(), today.isoformat())
    recent = c.execute(
        "SELECT t.*, a.name acc, a2.name acc2, c.name cat FROM transactions t JOIN accounts a ON a.id=t.account_id "
        "LEFT JOIN accounts a2 ON a2.id=t.to_account_id LEFT JOIN categories c ON c.id=t.category_id "
        "WHERE t.deleted_at IS NULL ORDER BY t.occurred_at DESC, t.id DESC LIMIT 12").fetchall()
    uncategorised = c.execute(
        "SELECT COUNT(*) FROM transactions t JOIN categories c ON c.id=t.category_id WHERE t.deleted_at IS NULL "
        "AND t.user_edited=0 AND (c.auto=1 OR c.name IN ('Інші доходи','Інші витрати'))").fetchone()[0]
    return render(request, "dashboard.html", u, categories=lists(c)["categories"], bal=b, day=day, month=month, recent=recent,
                  takings=takings, last_sync=get_setting(c, "last_sync_all"), review=uncategorised,
                  in_window=sync.in_work_window())


_bg = {"running": False}


def _run_sync_bg():
    try:
        sync.sync_all(get_con(), force=True)
    except Exception:
        log.exception("manual sync failed")
    finally:
        _bg["running"] = False


@app.post("/sync")
def sync_now(request: Request, u=Depends(current_user)):
    """Runs in the background: a full history refresh can take minutes, longer than a web request may live."""
    import threading
    if _bg["running"]:
        flash(request, "Оновлення вже виконується, зачекайте 1–3 хвилини і оновіть сторінку")
    else:
        _bg["running"] = True
        threading.Thread(target=_run_sync_bg, daemon=True).start()
        flash(request, "Оновлення запущено. Оновіть сторінку через хвилину-дві: дані з банків і Alteg підтягнуться.")
    return RedirectResponse(request.headers.get("referer", "/"), status_code=303)


# ---------- transactions ----------
def parse_tx_filters(qp) -> dict:
    """Query string -> filter dict. Multi-value: account, category, kind, source."""
    def many(name):
        return [v for v in qp.getlist(name) if v]
    def kop(name):
        try:
            return round(float(qp.get(name, "").replace(" ", "").replace(",", ".")) * 100) if qp.get(name) else None
        except ValueError:
            return None
    f = dict(accounts=[int(x) for x in many("account") if x.isdigit()],
             categories=[x for x in many("category") if x.isdigit() or x == "none"],
             kinds=[x for x in many("kind") if x in KINDS], sources=[x for x in many("source") if x in SRC],
             q=(qp.get("q") or "").strip(), date_from=qp.get("date_from") or "", date_to=qp.get("date_to") or "",
             amount_from=kop("amount_from"), amount_to=kop("amount_to"), deleted=qp.get("show") == "deleted",
             sort=qp.get("sort") if qp.get("sort") in ("date_desc", "date_asc", "amount_desc", "amount_asc") else "date_desc")
    return f


SORTS = {"date_desc": "t.occurred_at DESC, t.id DESC", "date_asc": "t.occurred_at ASC, t.id ASC",
         "amount_desc": "t.amount DESC, t.id DESC", "amount_asc": "t.amount ASC, t.id ASC"}


def tx_query(f: dict, order=True):
    where, args = ["t.deleted_at IS NOT NULL" if f["deleted"] else "t.deleted_at IS NULL"], []
    if f["date_from"]:
        where.append("t.occurred_at >= ?"); args.append(f["date_from"] + "T00:00:00")
    if f["date_to"]:
        where.append("t.occurred_at <= ?"); args.append(f["date_to"] + "T23:59:59")
    if f["accounts"]:
        ph = ",".join("?" * len(f["accounts"]))
        where.append(f"(t.account_id IN ({ph}) OR t.to_account_id IN ({ph}))"); args += f["accounts"] * 2
    if f["categories"]:
        ids = [int(x) for x in f["categories"] if x != "none"]
        parts = []
        if ids:
            parts.append(f"t.category_id IN ({','.join('?' * len(ids))})"); args += ids
        if "none" in f["categories"]:
            parts.append("(t.category_id IS NULL AND t.kind!='transfer')")
        where.append("(" + " OR ".join(parts) + ")")
    if f["kinds"]:
        where.append(f"t.kind IN ({','.join('?' * len(f['kinds']))})"); args += f["kinds"]
    if f["sources"]:
        where.append(f"t.source IN ({','.join('?' * len(f['sources']))})"); args += f["sources"]
    for word in f["q"].lower().split():   # every word must be found in the description or the counterparty
        where.append("(lower(t.description) LIKE ? OR lower(t.counterparty) LIKE ?)"); args += [f"%{word}%"] * 2
    if f["amount_from"] is not None:
        where.append("t.amount >= ?"); args.append(f["amount_from"])
    if f["amount_to"] is not None:
        where.append("t.amount <= ?"); args.append(f["amount_to"])
    sql = ("SELECT t.*, a.name acc, a2.name acc2, c.name cat, u.name author FROM transactions t "
           "JOIN accounts a ON a.id=t.account_id LEFT JOIN accounts a2 ON a2.id=t.to_account_id "
           "LEFT JOIN categories c ON c.id=t.category_id LEFT JOIN users u ON u.id=t.created_by "
           "WHERE " + " AND ".join(where))
    if order:
        sql += " ORDER BY " + SORTS[f["sort"]]
    return sql, args


def _items(request):
    return [(k, v) for k, v in request.query_params.multi_items() if k != "page" and v != ""]


def fu(items, set=None, drop=None, add=None) -> str:
    """Builds a query string from the current filters: set (replace key), drop ((key, value) pairs or key), add (append value)."""
    from urllib.parse import urlencode
    out = list(items)
    for k, v in (drop or []):
        out = [(a, b) for a, b in out if not (a == k and (v is None or b == v))]
    for k, v in (set or {}).items():
        out = [(a, b) for a, b in out if a != k]
        if v not in (None, ""):
            out.append((k, v))
    for k, v in (add or []):
        out.append((k, v))
    return "?" + urlencode(out)


def hl(text, words):
    """Escapes text and highlights searched words."""
    from markupsafe import Markup, escape
    t = str(escape(text or ""))
    for w in sorted(set(words), key=len, reverse=True):
        if w:
            t = re.sub("(" + re.escape(str(escape(w))) + ")", r"<mark>\1</mark>", t, flags=re.I)
    return Markup(t)


tpl.env.globals["fu"] = fu
tpl.env.globals["hl"] = hl


def lists(c):
    return dict(accounts=c.execute("SELECT * FROM accounts WHERE active=1 AND kind!='credit' ORDER BY sort,id").fetchall(),
                categories=c.execute("SELECT * FROM categories WHERE archived=0 ORDER BY kind DESC, name").fetchall())


def _tx_presets():
    today = datetime.now(config.KYIV).date()
    first = today.replace(day=1)
    pl_ = first - timedelta(days=1)
    return [("today", "Сьогодні", today, today), ("yesterday", "Вчора", today - timedelta(days=1), today - timedelta(days=1)),
            ("week", "7 днів", today - timedelta(days=6), today), ("month", "Цей місяць", first, today),
            ("prev_month", "Минулий місяць", pl_.replace(day=1), pl_)]


@app.get("/transactions", response_class=HTMLResponse)
def tx_list(request: Request, u=Depends(current_user)):
    c = get_con()
    qp = request.query_params
    f = parse_tx_filters(qp)
    page = max(int(qp.get("page", 1) or 1), 1) if (qp.get("page") or "1").isdigit() else 1
    sql, args = tx_query(f)
    total = c.execute(f"SELECT COUNT(*) FROM ({sql})", args).fetchone()[0]
    rows = c.execute(sql + " LIMIT 100 OFFSET ?", (*args, (page - 1) * 100)).fetchall()
    sums = {r["kind"]: (r["s"], r["n"]) for r in c.execute(f"SELECT kind, SUM(amount) s, COUNT(*) n FROM ({sql}) GROUP BY kind", args)}
    L = lists(c)
    items = _items(request)
    accn = {a["id"]: a["name"] for a in L["accounts"]}
    catn = {str(k["id"]): k["name"] for k in L["categories"]}
    chips = []
    for v in f["accounts"]:
        chips.append(("Рахунок: " + accn.get(v, str(v)), fu(items, drop=[("account", str(v))])))
    for v in f["categories"]:
        chips.append(("Категорія: " + ("без категорії" if v == "none" else catn.get(v, v)), fu(items, drop=[("category", v)])))
    for v in f["kinds"]:
        chips.append(("Тип: " + KINDS[v].lower(), fu(items, drop=[("kind", v)])))
    for v in f["sources"]:
        chips.append(("Джерело: " + SRC[v], fu(items, drop=[("source", v)])))
    if f["q"]:
        chips.append(("Опис: " + f["q"], fu(items, drop=[("q", None)])))
    if f["date_from"] or f["date_to"]:
        chips.append((f"Період: {f['date_from'] or '…'} — {f['date_to'] or '…'}", fu(items, drop=[("date_from", None), ("date_to", None)])))
    if f["amount_from"] is not None or f["amount_to"] is not None:
        chips.append((f"Сума: {qp.get('amount_from') or '0'} — {qp.get('amount_to') or '∞'}", fu(items, drop=[("amount_from", None), ("amount_to", None)])))
    if f["deleted"]:
        chips.append(("Видалені", fu(items, drop=[("show", None)])))
    # type counts for the segmented control (same filters but without the type)
    f2 = dict(f, kinds=[])
    sql2, args2 = tx_query(f2, order=False)
    kind_counts = {r["kind"]: r["n"] for r in c.execute(f"SELECT kind, COUNT(*) n FROM ({sql2}) GROUP BY kind", args2)}
    return render(request, "transactions.html", u, rows=rows, f=f, page=page, total=total, sums=sums, chips=chips, items=items,
                  kind_counts=kind_counts, presets=_tx_presets(), words=f["q"].lower().split(), **L,
                  qs=fu(items).lstrip("?"), SORT_LABELS={"date_desc": "спочатку нові", "date_asc": "спочатку старі",
                                                         "amount_desc": "сума ↓", "amount_asc": "сума ↑"})


@app.get("/transactions/export.csv")
def tx_csv(request: Request, u=Depends(current_user)):
    sql, args = tx_query(parse_tx_filters(request.query_params))
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["Дата", "Рахунок", "Тип", "Категорія", "Сума", "Контрагент", "Опис", "Джерело"])
    for t in get_con().execute(sql, args):
        w.writerow([t["occurred_at"].replace("T", " "), t["acc"] + (f" → {t['acc2']}" if t["acc2"] else ""),
                    KINDS[t["kind"]], t["cat"] or "", f"{t['amount'] / 100:.2f}".replace(".", ","),
                    t["counterparty"], t["description"], SRC.get(t["source"], t["source"])])
    return Response("\ufeff" + out.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=transactions.csv"})


@app.get("/transactions/export.xlsx")
def tx_xlsx(request: Request, u=Depends(current_user)):
    from openpyxl import Workbook
    sql, args = tx_query(parse_tx_filters(request.query_params))
    wb = Workbook()
    ws = wb.active
    ws.title = "Операції"
    ws.append(["Дата", "Рахунок", "Тип", "Категорія", "Сума", "Контрагент", "Опис", "Джерело"])
    for t in get_con().execute(sql, args):
        sign = -1 if t["kind"] == "expense" else 1
        ws.append([t["occurred_at"].replace("T", " "), t["acc"] + (f" → {t['acc2']}" if t["acc2"] else ""), KINDS[t["kind"]],
                   t["cat"] or "", sign * t["amount"] / 100, t["counterparty"], t["description"], SRC.get(t["source"], t["source"])])
    for col, wd in zip("ABCDEFGH", (18, 22, 14, 26, 14, 28, 60, 14)):
        ws.column_dimensions[col].width = wd
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=transactions.xlsx"})


def _tx_form_data(kind, account_id, to_account_id, amount, occurred_at, description, counterparty, category_id):
    if kind not in KINDS:
        raise HTTPException(400, "Невірний тип")
    d = dict(kind=kind, account_id=account_id, amount=parse_amount(amount), occurred_at=parse_dt(occurred_at),
             description=description.strip(), counterparty=counterparty.strip())
    if kind == "transfer":
        if not to_account_id or to_account_id == account_id:
            raise HTTPException(400, "Оберіть інший рахунок призначення")
        d.update(to_account_id=to_account_id, category_id=None)
    else:
        d.update(to_account_id=None, category_id=category_id or cat.categorize(
            get_con(), kind, description, counterparty))
    return d


@app.get("/transactions/new", response_class=HTMLResponse)
def tx_new(request: Request, kind: str = "expense", from_acc: int | None = None, to_acc: int | None = None,
           desc: str = "", u=Depends(current_user)):
    c = get_con()
    ids = {r["name"]: r["id"] for r in c.execute("SELECT id,name FROM accounts WHERE active=1 AND kind!='credit'")}

    def find(word):
        return next((i for n, i in ids.items() if word in n.lower()), None)
    presets = [(lbl, f, t, d) for lbl, f, t, d in (
        ("Інкасація: каса → сейф", find("каса"), find("сейф"), "Інкасація"),
        ("Поповнення карти: каса → робочі гроші", find("каса"), find("робочі"), "Поповнення карти через банкомат"))
        if f and t]
    return render(request, "tx_form.html", u, t=None, kind=kind, history=[], presets=presets,
                  pre=dict(from_acc=from_acc, to_acc=to_acc, desc=desc),
                  default_dt=datetime.now(config.KYIV).strftime("%Y-%m-%dT%H:%M"), **lists(c))


@app.post("/transactions/new")
def tx_create(request: Request, kind: str = Form(...), account_id: int = Form(...),
              to_account_id: int | None = Form(None), amount: str = Form(...), occurred_at: str = Form(...),
              description: str = Form(""), counterparty: str = Form(""), category_id: int | None = Form(None),
              u=Depends(current_user)):
    c = get_con()
    d = _tx_form_data(kind, account_id, to_account_id, amount, occurred_at, description, counterparty, category_id)
    with lock():
        insert_tx(c, user_id=u["id"], source="manual", **d)
        c.commit()
    flash(request, "Операцію додано")
    return RedirectResponse("/transactions", status_code=303)


@app.get("/transactions/{tid}", response_class=HTMLResponse)
def tx_edit_page(request: Request, tid: int, u=Depends(current_user)):
    c = get_con()
    t = c.execute("SELECT * FROM transactions WHERE id=?", (tid,)).fetchone()
    if not t:
        raise HTTPException(404)
    hist = c.execute("SELECT a.*, u.name uname FROM audit a LEFT JOIN users u ON u.id=a.user_id "
                     "WHERE a.tx_id=? ORDER BY a.id DESC", (tid,)).fetchall()
    return render(request, "tx_form.html", u, t=t, kind=t["kind"], history=hist,
                  default_dt=t["occurred_at"][:16], **lists(c))


@app.post("/transactions/{tid}")
def tx_update(request: Request, tid: int, kind: str = Form(...), account_id: int = Form(...),
              to_account_id: int | None = Form(None), amount: str = Form(...), occurred_at: str = Form(...),
              description: str = Form(""), counterparty: str = Form(""), category_id: int | None = Form(None),
              make_rule: str = Form(""), rule_text: str = Form(""), u=Depends(current_user)):
    c = get_con()
    if not c.execute("SELECT 1 FROM transactions WHERE id=?", (tid,)).fetchone():
        raise HTTPException(404)
    d = _tx_form_data(kind, account_id, to_account_id, amount, occurred_at, description, counterparty, category_id)
    with lock():
        update_tx(c, tid, u["id"], **d)
        extra = ""
        if make_rule and rule_text.strip() and d.get("category_id"):
            cat.add_rule(c, d["category_id"], rule_text, kind if kind != "transfer" else "any", priority=10)
            n = cat.apply_rule_to_existing(c, d["category_id"], rule_text, kind, u["id"])
            extra = f" Правило «{rule_text.strip()}» створено, перерозподілено операцій: {n}."
        c.commit()
    flash(request, "Збережено." + extra)
    return RedirectResponse("/transactions", status_code=303)


@app.post("/transactions/{tid}/category")
def tx_set_category(tid: int, category_id: int = Form(...), rule_text: str = Form(""), u=Depends(current_user)):
    """Inline category change from lists; with rule_text it becomes a rule for all similar transactions."""
    c = get_con()
    t = c.execute("SELECT * FROM transactions WHERE id=?", (tid,)).fetchone()
    k = c.execute("SELECT * FROM categories WHERE id=?", (category_id,)).fetchone()
    if not t or not k or t["kind"] == "transfer" or k["kind"] != t["kind"]:
        raise HTTPException(400, "Категорія не підходить до цієї операції")
    n = 0
    with lock():
        update_tx(c, tid, u["id"], note="категорію змінено вручну", category_id=category_id)
        if rule_text.strip():
            cat.add_rule(c, category_id, rule_text, t["kind"], priority=10)
            n = cat.apply_rule_to_existing(c, category_id, rule_text, t["kind"], u["id"], override=True)
        c.commit()
    return {"ok": True, "similar": n}


@app.post("/transactions/{tid}/delete")
def tx_delete(request: Request, tid: int, u=Depends(current_user)):
    with lock():
        update_tx(get_con(), tid, u["id"], deleted_at=now_local())
        get_con().commit()
    flash(request, "Видалено (історію збережено, можна відновити)")
    return RedirectResponse("/transactions", status_code=303)


@app.post("/transactions/{tid}/restore")
def tx_restore(request: Request, tid: int, u=Depends(current_user)):
    with lock():
        update_tx(get_con(), tid, u["id"], deleted_at=None)
        get_con().commit()
    flash(request, "Відновлено")
    return RedirectResponse(f"/transactions/{tid}", status_code=303)


# ---------- credits ----------
@app.get("/credits", response_class=HTMLResponse)
def credits_page(request: Request, u=Depends(current_user)):
    c = get_con()
    ops = c.execute("SELECT o.*, a.name aname FROM credit_ops o JOIN accounts a ON a.id=o.account_id "
                    "ORDER BY o.occurred_at DESC, o.id DESC LIMIT 300").fetchall()
    repay = c.execute(
        "SELECT t.*, a.name acc FROM transactions t JOIN accounts a ON a.id=t.account_id JOIN categories k ON k.id=t.category_id "
        "WHERE t.deleted_at IS NULL AND k.name='Кредит' AND t.kind='expense' ORDER BY t.occurred_at DESC LIMIT 100").fetchall()
    return render(request, "credits.html", u, credit=sync.credit_summary(c), ops=ops, repay=repay)


# ---------- cash by day ----------
def _cash_period(request, default_preset: str | None = None):
    today = datetime.now(config.KYIV).date()
    q = request.query_params
    pr = q.get("preset")
    if not pr and default_preset and not q.get("date_from") and not q.get("date_to"):
        pr = default_preset
    first = today.replace(day=1)
    prev_last = first - timedelta(days=1)
    presets = {"today": (today, today), "yesterday": (today - timedelta(days=1),) * 2,
               "7": (today - timedelta(days=6), today), "30": (today - timedelta(days=29), today),
               "month": (first, today), "prev_month": (prev_last.replace(day=1), prev_last)}
    if pr in presets:
        a, b = presets[pr]
        return a.isoformat(), b.isoformat(), pr
    a = q.get("date_from") or (today - timedelta(days=29)).isoformat()
    b = q.get("date_to") or today.isoformat()
    try:
        da, db = date.fromisoformat(a), date.fromisoformat(b)
    except ValueError:
        raise HTTPException(400, "Некоректна дата")
    if db < da or (db - da).days > 400:
        raise HTTPException(400, "Період має бути від 1 до 400 днів")
    return a, b, ("30" if not q.get("date_from") and not q.get("date_to") else "")


@app.get("/cash", response_class=HTMLResponse)
def cash_page(request: Request, u=Depends(current_user)):
    a, b, pr = _cash_period(request)
    return render(request, "cash.html", u, cash=reports.daily_cash(get_con(), a, b), d_from=a, d_to=b, preset=pr)


@app.get("/cash/detail", response_class=HTMLResponse)
def cash_detail(request: Request, u=Depends(current_user)):
    a, b, pr = _cash_period(request, "today")
    return render(request, "cash_detail.html", u, det=reports.takings_detail(get_con(), a, b), d_from=a, d_to=b, preset=pr,
                  pay_error=get_setting(get_con(), "pay_sync_error", ""))


@app.get("/api/cash/detail")
def api_cash_detail(request: Request, _=Depends(api_user)):
    a, b, _p = _cash_period(request, "today")
    d = reports.takings_detail(get_con(), a, b)
    k = lambda v: v / 100
    return JSONResponse({"from": a, "to": b, "total": k(d["total"]), "cash": k(d["cash"]), "cashless": k(d["cashless"]),
                         "payments": d["n"],
                         "masters": [{"name": m["name"], "role": m["role"], "total": k(m["total"]), "cash": k(m["cash"]),
                                      "cashless": k(m["cashless"]), "payments": m["n"],
                                      "services": [{"direction": x["category"], "service": x["title"], "count": x["n"],
                                                    "total": k(x["total"])} for x in m["services"]]} for m in d["masters"]],
                         "directions": [{"direction": c["title"], "total": k(c["total"]), "payments": c["n"]} for c in d["categories"]],
                         "services": [{"direction": x["category"], "service": x["title"], "count": x["n"], "total": k(x["total"])}
                                      for x in d["services"]]})


@app.get("/cash/export.xlsx")
def cash_xlsx(request: Request, u=Depends(current_user)):
    from openpyxl import Workbook
    a, b, _ = _cash_period(request)
    rep = reports.daily_cash(get_con(), a, b)
    wb = Workbook()
    ws = wb.active
    ws.title = "Каса по днях"
    ws.append(["Дата", "Готівка (каса)", "Безготівка (ФОП)", "Інше", "Виручка разом", "Витрати з каси", "Інкасація", "Внесено в касу", "Залишок каси на кінець дня"])
    for d in rep["days"][::-1]:
        ws.append([d["date"]] + [reports.uah(d[k]) for k in ("cash", "cashless", "other", "total", "expense", "incasso", "deposit", "closing")])
    t = rep["totals"]
    ws.append(["Разом"] + [reports.uah(t[k]) for k in ("cash", "cashless", "other", "total", "expense", "incasso", "deposit")])
    for col in "ABCDEFGHI":
        ws.column_dimensions[col].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=cash_{a}_{b}.xlsx"})


# ---------- P&L ----------
def _pl_period(request):
    today = datetime.now(config.KYIV).date()
    q = request.query_params
    first = today.replace(day=1)
    prev_last = first - timedelta(days=1)
    qs = date(today.year, 3 * ((today.month - 1) // 3) + 1, 1)
    presets = {"month": (first, today), "prev_month": (prev_last.replace(day=1), prev_last), "quarter": (qs, today),
               "all": (date.fromisoformat(config.HISTORY_START), today)}
    pr = q.get("preset") or ("prev_month" if not q.get("date_from") and not q.get("date_to") else "")
    if pr in presets:
        a, b = presets[pr]
        return a.isoformat(), b.isoformat(), pr
    a = q.get("date_from") or first.isoformat()
    b = q.get("date_to") or today.isoformat()
    try:
        da, db = date.fromisoformat(a), date.fromisoformat(b)
    except ValueError:
        raise HTTPException(400, "Некоректна дата")
    if db < da or (db - da).days > 800:
        raise HTTPException(400, "Період має бути від 1 до 800 днів")
    return a, b, ""


@app.get("/pl", response_class=HTMLResponse)
def pl_page(request: Request, u=Depends(current_user)):
    a, b, pr = _pl_period(request)
    acc = request.query_params.get("accrual", "1") == "1"
    return render(request, "pl.html", u, rep=pl.build(get_con(), a, b, acc), d_from=a, d_to=b, preset=pr, accrual=acc,
                  pay_error=get_setting(get_con(), "pay_sync_error", ""),
                  closed=get_con().execute("SELECT * FROM pl_closed ORDER BY month DESC").fetchall(), payroll_months=payroll.covered_months(get_con()), payroll_at=get_setting(get_con(), "payroll_synced_at"),
                  payroll_err=get_setting(get_con(), "payroll_error", ""), can_payroll=u["role"] in ("owner", "accountant"))


def _payroll_month(request):
    q = request.query_params.get("month")
    today = datetime.now(config.KYIV).date()
    if q and len(q) == 7:
        return q
    return (today.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")


@app.get("/payroll", response_class=HTMLResponse)
def payroll_page(request: Request, u=Depends(current_user)):
    c = get_con()
    ym = _payroll_month(request)
    y, m = map(int, ym.split("-"))
    ms = date(y, m, 1)
    nxt = (ms.replace(day=28) + timedelta(days=4)).replace(day=1)
    plan, rec = payroll.payment_plan(c, ym), payroll.payments_recorded(c, ym)
    costs = payroll.month_costs(c, ym)
    acc = c.execute("SELECT id,name FROM accounts WHERE kind='cash' AND active=1 ORDER BY sort,id").fetchall()
    lo, hi = ms, nxt - timedelta(days=1)
    esv = pl._esv_sum(c, lo, hi)
    unrecorded = max(plan["cash"] - rec["cash"], 0)
    months = sorted(payroll.covered_months(c), reverse=True)
    return render(request, "payroll.html", u, ym=ym, plan=plan, rec=rec, costs=costs, esv=esv, unrecorded=unrecorded, months=months,
                  accounts=acc, pay_date=(nxt + timedelta(days=6)).isoformat(), can_record=u["role"] in ("owner", "accountant"),
                  order=payroll.SECTION_ORDER, month_name=ms.strftime("%m.%Y"))


@app.post("/payroll/record-cash")
async def payroll_record_cash(request: Request, u=Depends(staff_cat)):
    """Books the cash part of a month's payroll (paid on the 7th from the safe / cash desk). Today's real balances already
    include this payout, so the opening balance is raised by the same amount: balances stay as they are, history gets right."""
    form = await request.form()
    month, pay_date = form.get("month", ""), form.get("pay_date", "")
    c = get_con()
    cat_id = cat.get_or_create_category(c, "Зарплата", "expense")
    marker = f"Виплата ЗП за {month} готівкою (за зарплатною таблицею)"
    done, skipped = 0, 0
    with lock():
        for key, val in form.multi_items():
            if not key.startswith("amt_") or not str(val).strip():
                continue
            aid = int(key[4:])
            acc = c.execute("SELECT * FROM accounts WHERE id=? AND kind='cash'", (aid,)).fetchone()
            if not acc:
                continue
            if c.execute("SELECT 1 FROM transactions WHERE account_id=? AND description=? AND deleted_at IS NULL", (aid, marker)).fetchone():
                skipped += 1
                continue
            amt = parse_amount(str(val))
            insert_tx(c, account_id=aid, kind="expense", amount=amt, occurred_at=parse_dt(pay_date + "T10:00"), description=marker,
                      category_id=cat_id, source="manual", user_id=u["id"])
            c.execute("UPDATE accounts SET opening_balance=opening_balance+? WHERE id=?", (amt, aid))
            done += 1
        c.commit()
    flash(request, f"Записано виплат: {done}" + (f"; пропущено (вже записано): {skipped}" if skipped else "") + ". Поточні залишки не змінились.")
    return RedirectResponse(f"/payroll?month={month}", status_code=303)


@app.post("/payroll/sync")
def payroll_sync(request: Request, u=Depends(staff_cat)):
    try:
        months = payroll.sync(get_con())
        flash(request, "Зарплатну таблицю оновлено, місяці: " + ", ".join(months))
    except Exception as e:
        get_con().rollback()
        flash(request, f"Не вдалось оновити таблицю: {e}")
    return RedirectResponse(request.headers.get("referer", "/pl"), status_code=303)


@app.get("/pl/export.xlsx")
def pl_xlsx(request: Request, u=Depends(current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    a, b, _ = _pl_period(request)
    rep = pl.build(get_con(), a, b, request.query_params.get("accrual", "1") == "1")
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    ws.append([f"P&L {a} — {b}" + (" (за місяцем нарахування)" if rep["accrual"] else " (за датою оплати)")])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Стаття", "Сума, ₴", "% від виручки", f"Попередній період {rep['prev_from']} — {rep['prev_to']}", "Зміна, %"])
    for l in rep["lines"]:
        ws.append([("    " * l["indent"]) + l["label"], None if l["amount"] is None else reports.uah(l["amount"]), l["pct"],
                   None if l["prev"] is None else reports.uah(l["prev"]), l["delta"]])
        if l["kind"] in ("head", "sub", "total", "grand"):
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
    ws.append([])
    ws.append(["Поза P&L (довідково)"])
    for label, v in rep["memo"]:
        ws.append([label, reports.uah(v)])
    if rep["months"]:
        w2 = wb.create_sheet("По місяцях")
        w2.append(["Місяць", "Виручка", "Матеріали", "Валовий прибуток", "Оплата праці", "Опер. витрати", "EBITDA", "Відсотки", "Податки", "Чистий прибуток"])
        for m in rep["months"]:
            w2.append([m["label"]] + [reports.uah(m[k]) for k in ("revenue", "cogs", "gross", "labour", "opex", "ebitda", "interest", "tax", "net")])
    for w in wb.worksheets:
        w.column_dimensions["A"].width = 58
        for col in "BCDEFGHIJ":
            w.column_dimensions[col].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=PnL_{a}_{b}.xlsx"})


@app.get("/api/pl")
def api_pl(request: Request, _=Depends(api_user)):
    a, b, _p = _pl_period(request)
    r = pl.build(get_con(), a, b, request.query_params.get("accrual", "1") == "1")
    k = lambda v: None if v is None else v / 100
    return JSONResponse({"from": a, "to": b, "accrual": r["accrual"], "kpis": {**r["kpis"], "avg_check": k(r["kpis"]["avg_check"]),
                         "rev_per_day": k(r["kpis"]["rev_per_day"]), "bep": k(r["kpis"]["bep"])},
                         "lines": [{"label": l["label"], "amount": k(l["amount"]), "pct": l["pct"], "prev": k(l["prev"]),
                                    "delta_pct": l["delta"], "kind": l["kind"], "indent": l["indent"]} for l in r["lines"]],
                         "memo": [{"label": a_, "amount": k(v)} for a_, v in r["memo"]],
                         "months": [{x: (y if x == "label" else y / 100) for x, y in m.items()} for m in r["months"]]})


# ---------- reports ----------
def _period(request):
    today = datetime.now(config.KYIV).date()
    q = request.query_params
    pr = q.get("preset")
    if pr in reports.presets(today):
        a, b = reports.presets(today)[pr]
        return a.isoformat(), b.isoformat(), pr
    a = q.get("date_from") or today.replace(day=1).isoformat()
    b = q.get("date_to") or today.isoformat()
    for s in (a, b):
        try:
            date.fromisoformat(s)
        except ValueError:
            raise HTTPException(400, "Некоректна дата")
    return a, b, ""


@app.get("/reports", response_class=HTMLResponse)
def reports_page(request: Request, u=Depends(current_user)):
    a, b, pr = _period(request)
    acc = int(request.query_params.get("account") or 0) or None
    c = get_con()
    return render(request, "reports.html", u, rep=reports.build(c, a, b, acc), d_from=a, d_to=b, preset=pr,
                  acc=acc, **lists(c))


@app.get("/reports/export.xlsx")
def reports_xlsx(request: Request, u=Depends(current_user)):
    a, b, _ = _period(request)
    acc = int(request.query_params.get("account") or 0) or None
    c = get_con()
    data = reports.to_xlsx(reports.build(c, a, b, acc), reports.balances(c))
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=report_{a}_{b}.xlsx"})


# ---------- categories ----------
@app.get("/categories", response_class=HTMLResponse)
def categories_page(request: Request, u=Depends(current_user)):
    c = get_con()
    cats = c.execute("SELECT c.*, (SELECT COUNT(*) FROM transactions t WHERE t.category_id=c.id AND t.deleted_at IS NULL) n "
                     "FROM categories c WHERE archived=0 ORDER BY kind DESC, auto DESC, name").fetchall()
    rules = c.execute("SELECT r.*, c.name cname FROM rules r JOIN categories c ON c.id=r.category_id "
                      "ORDER BY c.name, r.pattern").fetchall()
    return render(request, "categories.html", u, cats=cats, rules=rules, can_edit=u["role"] in ("owner", "accountant"))


@app.post("/categories")
def category_add(request: Request, name: str = Form(...), kind: str = Form(...), u=Depends(staff_cat)):
    if kind not in ("income", "expense") or not name.strip():
        raise HTTPException(400)
    cat.get_or_create_category(get_con(), name, kind)
    get_con().commit()
    flash(request, "Категорію додано")
    return RedirectResponse("/categories", status_code=303)


@app.post("/categories/{cid}/pl")
def category_pl(request: Request, cid: int, pl_group: str = Form(""), u=Depends(staff_cat)):
    if pl_group and pl_group not in pl.GROUP_NAMES:
        raise HTTPException(400)
    get_con().execute("UPDATE categories SET pl_group=? WHERE id=?", (pl_group, cid))
    get_con().commit()
    flash(request, "Стаття P&L збережена")
    return RedirectResponse("/categories", status_code=303)


@app.post("/categories/{cid}/exclude")
def category_exclude(request: Request, cid: int, exclude: str = Form("0"), u=Depends(staff_cat)):
    get_con().execute("UPDATE categories SET exclude=? WHERE id=?", (1 if exclude == "1" else 0, cid))
    get_con().commit()
    flash(request, "Збережено")
    return RedirectResponse("/categories", status_code=303)


@app.post("/categories/{cid}")
def category_edit(request: Request, cid: int, name: str = Form(""), merge_into: int | None = Form(None),
                  u=Depends(staff_cat)):
    c = get_con()
    with lock():
        if merge_into and merge_into != cid:
            n = cat.merge_category(c, cid, merge_into, u["id"])
            flash(request, f"Об'єднано, перенесено операцій: {n}")
        elif name.strip():
            c.execute("UPDATE categories SET name=?, auto=0 WHERE id=?", (name.strip(), cid))
            flash(request, "Збережено")
        c.commit()
    return RedirectResponse("/categories", status_code=303)


@app.post("/rules")
def rule_add(request: Request, pattern: str = Form(...), category_id: int = Form(...), direction: str = Form("any"),
             apply: str = Form(""), u=Depends(staff_cat)):
    c = get_con()
    if direction not in ("any", "income", "expense"):
        raise HTTPException(400)
    with lock():
        cat.add_rule(c, category_id, pattern, direction, priority=10)
        n = cat.apply_rule_to_existing(c, category_id, pattern, None if direction == "any" else direction,
                                       u["id"]) if apply else 0
        c.commit()
    flash(request, f"Правило додано. Перерозподілено операцій: {n}")
    return RedirectResponse("/categories", status_code=303)


@app.post("/rules/{rid}/delete")
def rule_del(request: Request, rid: int, u=Depends(staff_cat)):
    get_con().execute("DELETE FROM rules WHERE id=?", (rid,))
    get_con().commit()
    return RedirectResponse("/categories", status_code=303)


# ---------- accounts (owner) ----------
@app.get("/accounts", response_class=HTMLResponse)
def accounts_page(request: Request, u=Depends(current_user)):
    c = get_con()
    bal = account_balances(c)
    accs = c.execute("SELECT * FROM accounts ORDER BY active DESC, sort, id").fetchall()
    logs = c.execute("SELECT l.*, a.name aname FROM sync_log l LEFT JOIN accounts a ON a.id=l.account_id "
                     "ORDER BY l.id DESC LIMIT 15").fetchall()
    return render(request, "accounts.html", u, accs=accs, bal=bal, logs=logs, is_owner=u["role"] == "owner",
                  last_sync=get_setting(c, "last_sync_all"))


@app.post("/accounts")
def account_save(request: Request, id: int | None = Form(None), name: str = Form(...), kind: str = Form(...),
                 source: str = Form(...), iban: str = Form(""), token: str = Form(""), opening: str = Form(""),
                 current: str = Form(""),
                 active: str = Form("1"), u=Depends(owner_only)):
    if kind not in ACC_KINDS or source not in SRC:
        raise HTTPException(400)
    c = get_con()
    def money_in(v):
        try:
            return round(float(v.replace(" ", "").replace("\u00a0", "").replace(",", ".")) * 100)
        except ValueError:
            raise HTTPException(400, "Некоректна сума залишку")

    op = money_in(opening) if opening.strip() else None
    cur_target = money_in(current) if current.strip() else None
    with lock():
        if id:
            c.execute("UPDATE accounts SET name=?,kind=?,source=?,iban=?,active=? WHERE id=?",
                      (name.strip(), kind, source, iban.strip().replace(" ", ""), int(active), id))
            if token.strip():
                c.execute("UPDATE accounts SET token_enc=? WHERE id=?", (encrypt(token.strip()), id))
            if op is not None:
                c.execute("UPDATE accounts SET opening_balance=?, opening_auto=1 WHERE id=?", (op, id))
            if cur_target is not None:  # balance as of now: opening = target - everything booked so far
                booked = account_balances(c)[id] - c.execute("SELECT opening_balance FROM accounts WHERE id=?", (id,)).fetchone()[0]
                c.execute("UPDATE accounts SET opening_balance=?, opening_auto=1 WHERE id=?", (cur_target - booked, id))
        else:
            c.execute("INSERT INTO accounts(name,kind,source,iban,token_enc,opening_balance,sort) "
                      "VALUES(?,?,?,?,?,?,(SELECT COALESCE(MAX(sort),0)+1 FROM accounts))",
                      (name.strip(), kind, source, iban.strip().replace(" ", ""),
                       encrypt(token.strip()), cur_target if cur_target is not None else (op or 0)))
        c.commit()
    flash(request, "Рахунок збережено")
    return RedirectResponse("/accounts", status_code=303)


@app.post("/accounts/{aid}/test")
def account_test(request: Request, aid: int, u=Depends(owner_only)):
    from .security import decrypt
    a = get_con().execute("SELECT * FROM accounts WHERE id=?", (aid,)).fetchone()
    try:
        if a["source"] != "privat":
            raise privat.PrivatError("Перевірка лише для рахунків Приватбанку")
        bal = privat.fetch_balance(decrypt(a["token_enc"]), a["iban"])
        flash(request, f"З'єднання з банком працює. Залишок: {money(bal)}")
    except Exception as e:
        flash(request, f"Помилка: {e}")
    return RedirectResponse("/accounts", status_code=303)


# ---------- users (owner) ----------
@app.get("/users", response_class=HTMLResponse)
def users_page(request: Request, u=Depends(owner_only)):
    return render(request, "users.html", u, users=get_con().execute("SELECT * FROM users ORDER BY id").fetchall())


@app.post("/users")
def user_save(request: Request, email: str = Form(...), name: str = Form(...), role: str = Form(...),
              password: str = Form(...), u=Depends(owner_only)):
    if role not in ROLES or len(password) < 8:
        raise HTTPException(400, "Пароль мінімум 8 символів")
    c = get_con()
    try:
        c.execute("INSERT INTO users(email,name,role,pw_hash) VALUES(?,?,?,?)",
                  (email.strip().lower(), name.strip(), role, hash_password(password)))
        c.commit()
        flash(request, "Користувача додано")
    except Exception:
        flash(request, "Такий email вже існує")
    return RedirectResponse("/users", status_code=303)


@app.post("/users/{uid}")
def user_update(request: Request, uid: int, role: str = Form(...), active: str = Form("1"),
                password: str = Form(""), u=Depends(owner_only)):
    if role not in ROLES:
        raise HTTPException(400)
    c = get_con()
    if uid == u["id"] and (role != "owner" or active != "1"):
        flash(request, "Не можна позбавити себе доступу власника")
        return RedirectResponse("/users", status_code=303)
    c.execute("UPDATE users SET role=?, active=? WHERE id=?", (role, int(active), uid))
    if password:
        if len(password) < 8:
            raise HTTPException(400, "Пароль мінімум 8 символів")
        c.execute("UPDATE users SET pw_hash=? WHERE id=?", (hash_password(password), uid))
    c.commit()
    flash(request, "Збережено")
    return RedirectResponse("/users", status_code=303)


@app.post("/password")
def change_password(request: Request, old: str = Form(...), new: str = Form(...), u=Depends(current_user)):
    if not verify_password(old, u["pw_hash"]) or len(new) < 8:
        flash(request, "Невірний поточний пароль або новий занадто короткий")
    else:
        get_con().execute("UPDATE users SET pw_hash=? WHERE id=?", (hash_password(new), u["id"]))
        get_con().commit()
        flash(request, "Пароль змінено")
    return RedirectResponse("/", status_code=303)


# ---------- JSON API (session or X-API-Token) ----------
@app.get("/api/balances")
def api_balances(_=Depends(api_user)):
    return {**reports.balances(get_con()), **sync.credit_summary(get_con())}


@app.get("/api/cash")
def api_cash(request: Request, _=Depends(api_user)):
    a, b, _p = _cash_period(request)
    r = reports.daily_cash(get_con(), a, b)
    return JSONResponse({**r, "days": [{k: (v / 100 if k != "date" else v) for k, v in d.items()} for d in r["days"]],
                         "totals": {k: v / 100 for k, v in r["totals"].items()}})


@app.get("/api/report")
def api_report(request: Request, _=Depends(api_user)):
    a, b, _p = _period(request)
    acc = int(request.query_params.get("account") or 0) or None
    r = reports.build(get_con(), a, b, acc)
    r["transactions"] = [dict(kind=t["kind"], at=t["occurred_at"], account=t["acc"], to=t["acc2"],
                              category=t["cat"], amount=t["amount"] / 100, counterparty=t["counterparty"],
                              description=t["description"], source=t["source"])
                         for t in r["transactions"]] if request.query_params.get("details") else []
    return JSONResponse(r)
